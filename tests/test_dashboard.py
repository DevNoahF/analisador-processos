"""Integração local real com PDF sintético; nenhuma chamada paga."""
import io
import os
from pathlib import Path
import tempfile
import time
import unittest
from unittest.mock import patch

TEST_DATA = tempfile.TemporaryDirectory()
os.environ['SHR_DATA_DIR'] = TEST_DATA.name
from fastapi.testclient import TestClient
import fitz
from openpyxl import load_workbook
from app.main import APP, DATA, LOCK, SESSIONS, get_job, put_job
from plan_workbook import SCHEMA

HEADERS = {'X-SHR-Client': 'dashboard'}


def pdf_bytes():
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 50), 'CAPA PROCESSO\nNº do processo 5000001-00.2025.8.21.0001\nDocumento sintetico para teste, sem autos reais.')
    page = doc.new_page()
    page.insert_text((50, 50), 'Peca ficticia de teste. Nao representa atuacao, substabelecimento ou decisao judicial.\n' * 3)
    page.insert_text((50, 780), 'Processo 5000001-00.2025.8.21.0001/RS, Evento 3, PET1, Página 1')
    return doc.tobytes()


class Dashboard(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.client = TestClient(APP, headers=HEADERS)
        r = cls.client.post('/api/login', json={'password': 'senha-ficticia-apenas-teste'})
        assert r.status_code == 200, r.text

    def make_job(self, mode='new'):
        result = self.client.post('/api/jobs', json={'name': 'LOTE SINTÉTICO — TESTE', 'mode': mode})
        self.assertEqual(result.status_code, 200, result.text)
        return result.json()['id']

    def extracted_job(self, mode='new'):
        jid = self.make_job(mode)
        r = self.client.post(f'/api/jobs/{jid}/files', files={'file': ('teste.pdf', pdf_bytes(), 'application/pdf')})
        self.assertEqual(r.status_code, 200, r.text)
        fid = r.json()['id']
        self.assertEqual(self.client.post(f'/api/jobs/{jid}/extract').status_code, 200)
        deadline = time.monotonic() + 10
        while time.monotonic() < deadline:
            job = self.client.get(f'/api/jobs/{jid}').json()
            if not job['busy']:
                self.assertIsNone(job['error'], job)
                return jid, fid, job
            time.sleep(.02)
        self.fail('Extração não terminou em dez segundos.')

    def test_authenticated_upload_extract_visual_and_export(self):
        jid, fid, job = self.extracted_job()
        self.assertEqual(job['documents'][0]['pages'], 2)
        self.assertEqual(job['records'][0]['process'], '5000001-00.2025.8.21.0001')
        self.assertEqual(job['decisions'][0]['status'], 'PENDENTE')
        image = self.client.get(f'/api/jobs/{jid}/images/{fid}/1')
        self.assertTrue(image.content.startswith(b'\x89PNG'))
        p = self.client.get(f'/api/jobs/{jid}/pages/{fid}/1').json()
        r = self.client.post(f'/api/jobs/{jid}/pages/{fid}/1', json={'text': p['text'], 'source': 'Conferencia sintetica pelo teste automatizado.'})
        self.assertEqual(r.status_code, 200)
        raw = job['records']
        raw[0]['adverse_party'] = '=HYPERLINK("http://example.com","injection")'
        r = self.client.put(f'/api/jobs/{jid}/records', json={'records': raw})
        self.assertEqual(r.status_code, 200, r.text)
        r = self.client.post(f'/api/jobs/{jid}/export')
        self.assertEqual(r.status_code, 200, r.text)
        book = load_workbook(io.BytesIO(self.client.get(r.json()['url']).content))
        sheet = book[SCHEMA['sheet']]
        self.assertEqual(sheet.max_column, 37)
        self.assertEqual(sheet['B2'].value, 'PENDENTE')
        self.assertEqual(sheet['H2'].data_type, 's')
        self.assertIsNone(sheet['L2'].value)
        self.assertEqual(sheet['U2'].value, 'Pendente')
        self.assertEqual(self.client.post(f'/api/jobs/{jid}/export').status_code, 200)
        report = self.client.get(f'/api/jobs/{jid}/report').json()
        self.assertEqual(report['new_rows'], 1)
        self.assertFalse(report['review']['coverage'][0]['text_review_complete'])

    def test_append_preserves_formula_and_blocks_duplicate(self):
        jid, _, job = self.extracted_job()
        raw = job['records']
        raw[0].update(tariff_key='2022_1_DEFESA', uf='RS')
        self.client.put(f'/api/jobs/{jid}/records', json={'records': raw})
        self.client.post(f'/api/jobs/{jid}/export')
        original = self.client.get(f'/api/jobs/{jid}/download').content
        book = load_workbook(io.BytesIO(original))
        book['FORÇA TAREFA 2']['S2'] = '=1+2'
        book['REGRAS']['F3'] = 12345
        output = io.BytesIO()
        book.save(output)
        second, _, job2 = self.extracted_job('append')
        r = self.client.post(f'/api/jobs/{second}/files', files={'file': ('original.xlsx', output.getvalue(), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})
        self.assertEqual(r.status_code, 200, r.text)
        raw2 = job2['records']
        raw2[0].update(tariff_key='2022_1_DEFESA', uf='RS')
        self.client.put(f'/api/jobs/{second}/records', json={'records': raw2})
        r = self.client.post(f'/api/jobs/{second}/export')
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(r.json()['conflicts'], 1)
        self.assertEqual(r.json()['rows'], 0)
        final = load_workbook(io.BytesIO(self.client.get(r.json()['url']).content), data_only=False)
        self.assertEqual(final['FORÇA TAREFA 2']['S2'].value, '=1+2')
        self.assertEqual(final['REGRAS']['F3'].value, 12345)

    def test_invalid_upload_auth_origin_and_credentials(self):
        stranger = TestClient(APP, headers=HEADERS)
        self.assertEqual(stranger.get('/api/jobs').status_code, 401)
        self.assertEqual(self.client.post('/api/jobs', headers={'Origin': 'https://example.com'}, json={'name':'x'}).status_code, 403)
        self.assertEqual(self.client.get('/api/jobs', headers={'Host': 'attacker.example'}).status_code, 403)
        jid = self.make_job()
        r = self.client.post(f'/api/jobs/{jid}/files', files={'file': ('fake.pdf', b'not a pdf', 'application/pdf')})
        self.assertEqual(r.status_code, 400)
        secret = 'secret-do-not-reflect'
        r = self.client.post(f'/api/jobs/{jid}/review', json={'api_key': secret})
        self.assertEqual(r.status_code, 422)
        self.assertNotIn(secret, r.text)
        self.assertEqual(self.client.get('/api/jobs/../../account.json').status_code, 404)

    def test_pending_reading_and_budget_are_visible(self):
        jid, fid, job = self.extracted_job()
        r = self.client.post(f'/api/jobs/{jid}/preflight', json={'provider':'gemini','model':'not-configured','max_calls':1})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertTrue(r.json()['exceeds_call_budget'])
        self.assertEqual(r.json()['unique_pdfs'], 1)
        r = self.client.post(f'/api/jobs/{jid}/review', json={'provider':'gemini','model':'not-configured','api_key':'fake-key-only-test'})
        self.assertEqual(r.status_code, 400)
        self.assertNotIn('fake-key-only-test', ''.join(p.read_text() for p in DATA.glob('**/*.json')))

    def test_shared_password_gives_multiple_visitors_same_workspace(self):
        from app import main
        original = main.DATA
        with tempfile.TemporaryDirectory() as folder, patch.dict(os.environ, {'SHR_SHARED_PASSWORD': 'senha-compartilhada-para-equipe'}):
            main.DATA = Path(folder)
            main.DATA.mkdir(mode=0o700, exist_ok=True)
            first = TestClient(APP, headers=HEADERS)
            second = TestClient(APP, headers=HEADERS)
            self.assertEqual(first.get('/api/session').json()['shared_password_configured'], True)
            self.assertEqual(first.post('/api/login', json={'password':'senha-errada-para-teste'}).status_code, 401)
            self.assertEqual(first.post('/api/login', json={'password':'senha-compartilhada-para-equipe'}).status_code, 200)
            self.assertEqual(second.post('/api/login', json={'password':'senha-compartilhada-para-equipe'}).status_code, 200)
            created = first.post('/api/jobs', json={'name':'Lote comum','mode':'new'}).json()
            self.assertEqual([job['name'] for job in second.get('/api/jobs').json()], ['Lote comum'])
            first.post('/api/logout')
            self.assertEqual([job['id'] for job in second.get('/api/jobs').json()], [created['id']])
            self.assertNotIn('senha-compartilhada-para-equipe', (main.DATA / 'account.json').read_text())
            main.DATA = original


if __name__ == '__main__':
    unittest.main()
