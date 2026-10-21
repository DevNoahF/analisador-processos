"""Exportação conservadora: cópia integral, literais e conflitos explícitos."""
from copy import copy
from datetime import date
from io import BytesIO
from pathlib import Path
import json
import zipfile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from plan_workbook import SCHEMA, plan_write
from decidir_lote import load_tariffs

GREEN = '173D31'


def validate_xlsx(path):
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
        if len(names) > 3000 or sum(i.file_size for i in z.infolist()) > 80_000_000:
            raise ValueError('Planilha excede o limite de expansão de 80 MB.')
        if any(i.file_size > 20_000_000 for i in z.infolist()):
            raise ValueError('Componente XLSX muito grande.')
        # Recusar recursos que o escritor não promete preservar.
        unsupported = ('vba', 'externallinks/', 'embeddings/', 'activex/', 'drawings/',
                       'pivot', 'slicer', 'connections.xml', 'customxml/', 'querytables/')
        if any(any(part in n.lower() for part in unsupported) for n in names):
            raise ValueError('Planilha com macros, links ou recursos complexos. Use criação de nova planilha para preservar o original.')
        if '[Content_Types].xml' not in names or 'xl/workbook.xml' not in names:
            raise ValueError('Arquivo não é uma planilha XLSX válida.')
    book = load_workbook(BytesIO(Path(path).read_bytes()), data_only=False, keep_links=False)
    if any(c.hyperlink and str(c.hyperlink.target or '').lower().startswith(('file:', 'javascript:'))
           for ws in book for row in ws for c in row):
        raise ValueError('Planilha contém links não permitidos.')
    return book


def literal(cell, value):
    cell.value = value
    if isinstance(value, str):
        cell.data_type = 's'


def header(ws, values, row=1):
    for i, value in enumerate(values, 1):
        cell = ws.cell(row, i)
        literal(cell, value)
        cell.fill = PatternFill('solid', fgColor=GREEN)
        cell.font = Font(color='FFFFFF', bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[row].height = 34
    ws.freeze_panes = f'A{row + 1}'


def export_workbook(raw, review, *, mode, existing=None):
    from zoneinfo import ZoneInfo
    from datetime import datetime
    today = datetime.now(ZoneInfo('America/Sao_Paulo')).date().isoformat()
    plan = plan_write(raw, review, mode=mode, review_date=today, existing=existing)
    book = validate_xlsx(existing) if existing else Workbook()
    if not existing:
        book.active.title = SCHEMA['sheet']
    ws = book[SCHEMA['sheet']]
    if not existing:
        header(ws, SCHEMA['headers'])
        for i in range(1, 38):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = 24 if i != 6 else 31
    else:
        for i, name in enumerate(SCHEMA['headers'], 1):
            if ws.cell(1, i).value is None:
                literal(ws.cell(1, i), name)
                ws.cell(1, i)._style = copy(ws.cell(1, 1)._style)
    for item in plan['append']:
        for i, value in enumerate(item['values'], 1):
            cell = ws.cell(item['excel_row'], i)
            if i in SCHEMA['date_columns'] and value:
                try:
                    value = date.fromisoformat(value)
                except (ValueError, TypeError):
                    pass
            literal(cell, value)
            if i in SCHEMA['date_columns']:
                cell.number_format = 'dd/mm/yyyy'
            if i in SCHEMA['money_columns']:
                cell.number_format = '"R$" #,##0.00'
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    if not existing:
        ws.auto_filter.ref = ws.dimensions
        rules = book.create_sheet('REGRAS')
        rules.merge_cells('A1:I1')
        rules['A1'] = 'Regras de enquadramento e remuneração'
        header(rules, ['Chave', 'Contrato', 'Tabela', 'Procedimento', 'Fase/hipótese', 'Valor', 'Autorização', 'Condição', 'Fonte'], 2)
        for tariff in load_tariffs().values():
            rules.append([tariff[k] for k in ['key', 'contract', 'table', 'category', 'act', 'amount_brl', 'authorization', 'conditions', 'source']])
        for col in 'ABCDEFGHI':
            rules.column_dimensions[col].width = 35
        rules.append(['REGRAS DE DECISÃO', 'RS/SC tem precedência territorial; demais UFs usam Subs e ano do CNJ.'])
        rules.append(['PENDÊNCIAS', 'Valores de 2025 e notas ausentes exigem confirmação documental. Não libera faturamento.'])
    else:
        # Nunca atualizar parâmetros históricos de REGRAS silenciosamente.
        plan['rules_preserved'] = True
    sup = book['SUPERVISAO'] if 'SUPERVISAO' in book else book.create_sheet('SUPERVISAO')
    if sup.max_row > 1 or sup['A1'].value:
        if [sup.cell(1, i).value for i in range(1, 11)] != SCHEMA['supervision_headers']:
            raise ValueError('Aba SUPERVISAO incompatível; exportação interrompida para preservar o modelo.')
    else:
        header(sup, SCHEMA['supervision_headers'])
    for row in plan['supervision_rows']:
        for i, value in enumerate(row, 1):
            literal(sup.cell(sup.max_row + 1 if i == 1 else target, i), value)
            if i == 1:
                target = sup.max_row
    report = book.create_sheet('COBERTURA')
    header(report, ['PDF', 'Páginas', 'Texto revisado', 'Leitura conferida', 'Pendências'])
    for doc in review.get('coverage', []):
        report.append([doc['file_name'], doc['page_count'], str(doc.get('text_review_complete', False)),
                       str(doc.get('reading_complete', False)), json.dumps(doc.get('reading_pending', []), ensure_ascii=False)])
    report.append(['Consumo', json.dumps(review.get('usage', {}), ensure_ascii=False)])
    report.append(['Conflitos', json.dumps(plan['conflicts'], ensure_ascii=False)])
    report.append(['REGRAS', 'Aba histórica preservada; comparar versão e valores antes de atualizar parâmetros.' if existing else 'Tarifas de referência de 21/10/2026.'])
    # Todo texto externo deve ser literal, inclusive metadados de arquivos.
    for row in report:
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = 's'
    result = BytesIO()
    book.save(result)
    result.seek(0)
    check = load_workbook(result, data_only=False)
    if existing:
        original = load_workbook(BytesIO(Path(existing).read_bytes()), data_only=False, keep_links=False)
        for sheet in original:
            for row in sheet:
                for cell in row:
                    if cell.value is not None and check[sheet.title][cell.coordinate].value != cell.value:
                        raise ValueError('Verificação detectou alteração de célula histórica; exportação bloqueada.')
    return result.getvalue(), plan
